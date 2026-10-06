import openpyxl
import pymysql
import re
from datetime import datetime

EXCEL_PATH = '/Users/jerry/Documents/Repositories/shadcn-admin/assets/database_structure/2025年11月30日 经营表.xlsx'
SHEET_NAME = '2025年备件'

DB_CONFIG = {
    'host': '101.43.54.73',
    'port': 3306,
    'database': 'jiehong',
    'user': 'jiehong',
    'password': 'snaketzy123$',
    'charset': 'utf8mb4'
}

COLUMN_MAPPING = {
    'case_invoice_number': '发票号',
    'case_order_amount': '订单金额',
    'case_owner_custom_packing_transport_fee': '报关,包装,运费（船东）',
    'case_invoice_amount': '发票金额',
    'case_bank_fee': '银行手续费',
    'case_actual_amount': '实际到账',
    'case_pay_to_supplier': '应付供应商',
    'case_supplier_custom_packing_transport_fee': '报关,包装,运费（供应商）',
    'case_payment_fee': '付款手续费',
    'case_payment_date': '支付日期',
    'case_deadline_for_owner': '船东应付款时间',
    'case_receipt_date': '到账日期',
    'case_settlement_remark': '备注',
}

DATETIME_FIELDS = {'case_payment_date', 'case_deadline_for_owner', 'case_receipt_date'}

def parse_datetime_value(val):
    if val is None or (isinstance(val, str) and val.strip() == ''):
        return None
    if isinstance(val, datetime):
        return val
    if isinstance(val, (int, float)):
        try:
            s = str(int(val))
            if len(s) == 8:
                return datetime.strptime(s, '%Y%m%d')
        except:
            pass
    if isinstance(val, str):
        s = val.strip()
        if not s:
            return None
        fmts = ['%Y-%m-%d %H:%M:%S', '%Y-%m-%d', '%Y/%m/%d', '%Y%m%d', '%Y/%m/%d %H:%M:%S']
        for f in fmts:
            try:
                return datetime.strptime(s, f)
            except:
                continue
    return None

def parse_numeric_value(val):
    if val is None:
        return None
    if isinstance(val, (int, float)):
        s = str(val)
        if len(s) > 11:
            return s[:11]
        return s
    if isinstance(val, str):
        s = val.strip()
        if not s:
            return None
        if s.startswith('#') or s in ('#REF!', '#VALUE!', '#DIV/0!', '#N/A', '#NAME?'):
            return None
        s = s.replace(',', '')
        try:
            float(s)
            if len(s) > 11:
                return s[:11]
            return s
        except:
            return None
    return None

def parse_string_value(val, max_len=11):
    if val is None:
        return None
    if isinstance(val, (int, float)):
        s = str(val)
        if len(s) > max_len:
            return s[:max_len]
        return s
    if isinstance(val, str):
        s = val.strip()
        if not s:
            return None
        if len(s) > max_len:
            return s[:max_len]
        return s
    return None

def main():
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb[SHEET_NAME]

    header_row = 1
    header_idx = {}
    for col in range(1, ws.max_column + 1):
        h = ws.cell(row=header_row, column=col).value
        if h is not None:
            header_idx[str(h).strip()] = col

    print(f'Excel 表头列映射: {header_idx}')

    rows_data = []
    skipped_empty = 0
    total_rows = 0

    for row_idx in range(header_row + 1, ws.max_row + 1):
        total_rows += 1
        row_dict = {}
        has_data = False

        for db_field, excel_header in COLUMN_MAPPING.items():
            col_idx = header_idx.get(excel_header)
            if col_idx is None:
                row_dict[db_field] = None
                continue
            cell_val = ws.cell(row=row_idx, column=col_idx).value

            if db_field in DATETIME_FIELDS:
                row_dict[db_field] = parse_datetime_value(cell_val)
            elif db_field == 'case_settlement_remark':
                row_dict[db_field] = parse_string_value(cell_val, 500)
            elif db_field == 'case_invoice_number':
                row_dict[db_field] = parse_string_value(cell_val, 100)
            else:
                row_dict[db_field] = parse_numeric_value(cell_val)

            if row_dict[db_field] is not None and str(row_dict[db_field]).strip() != '':
                has_data = True

        if not has_data:
            skipped_empty += 1
            continue

        row_dict['case_id'] = None
        row_dict['case_exchange_rate'] = None
        rows_data.append(row_dict)

    print(f'Excel 总行数(不含表头): {total_rows}')
    print(f'有效数据行数: {len(rows_data)}')
    print(f'跳过空行数: {skipped_empty}')

    if not rows_data:
        print('无有效数据，终止导入')
        return

    fields = list(COLUMN_MAPPING.keys()) + ['case_id', 'case_exchange_rate']
    placeholders = ', '.join(['%s'] * len(fields))
    field_list = ', '.join(fields)
    sql = f"INSERT INTO case_settlement_list ({field_list}) VALUES ({placeholders})"

    connection = None
    try:
        connection = pymysql.connect(**DB_CONFIG)
        with connection.cursor() as cursor:
            cursor.execute("SELECT COUNT(*) FROM case_settlement_list")
            before_count = cursor.fetchone()[0]
            print(f'导入前 case_settlement_list 行数: {before_count}')

            values_list = []
            for r in rows_data:
                values_list.append(tuple(r[f] for f in fields))

            cursor.executemany(sql, values_list)
            connection.commit()

            cursor.execute("SELECT COUNT(*) FROM case_settlement_list")
            after_count = cursor.fetchone()[0]
            inserted = after_count - before_count
            print(f'导入后 case_settlement_list 行数: {after_count}')
            print(f'成功插入行数: {inserted}')

            cursor.execute("SELECT * FROM case_settlement_list ORDER BY case_settlement_id DESC LIMIT 3")
            sample = cursor.fetchall()
            print('\n=== 最后3条插入样例 ===')
            col_names = [d[0] for d in cursor.description]
            for s in sample:
                row_display = {}
                for i, cn in enumerate(col_names):
                    row_display[cn] = s[i]
                print(row_display)

    except Exception as e:
        print(f'导入失败: {e}')
        import traceback
        traceback.print_exc()
        if connection:
            connection.rollback()
    finally:
        if connection:
            connection.close()

if __name__ == '__main__':
    main()
